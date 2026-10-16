"""Band-level trends (0-4k, 1-4k, 4-10k, 10-24k Hz) for every recording, condition classification, Excel workbook.

Spec: design/bdd-band-trends.md.   Usage: python band_trends.py <db_dir> <raw_dir> <out_dir>
Writes <out_dir>/cache/*.npz (per-frame band mean-square), <out_dir>/png/*.png, <out_dir>/MSM_頻帶趨勢與工況.xlsx.
No labels are read. VF-10 test sessions are reported in fixed 10-s segments, never by labelled interval.
"""
import os
import sqlite3
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill

from msm_io import find, read_wav_int16
from plot_db_envelope import rng  # 同時套用中文字型設定

NFFT, HOP_S = 4096, 0.1
BANDS = [("0-4k", 0, 4000), ("1-4k", 1000, 4000), ("4-10k", 4000, 10000), ("10-24k", 10000, 24001)]
PARTITION = ["0-4k", "4-10k", "10-24k"]  # 三段相加＝全頻
CH_NAME = {"s0": "s0 空氣麥克風", "s1": "s1 聽診器", "s2": "s2 聽診器"}
SHORT = {"vm20i": "VM20i", "vmx30ui": "VMX30Ui", "vf10": "VF-10"}
VF10_REF = [2, 7, 12]
SEG_S = 10.0
COLORS = {"0-4k": "#2b5c7a", "1-4k": "#7fb3d5", "4-10k": "#c47f1a", "10-24k": "#8a3b6e"}


def band_ms(x, sr):
    """Per-frame mean-square per band (Parseval on a Hann-windowed 4096-point FFT, one frame every 0.1 s)."""
    hop = int(sr * HOP_S)
    n = (len(x) - NFFT) // hop + 1
    win = np.hanning(NFFT).astype(np.float32)
    f = np.fft.rfftfreq(NFFT, 1 / sr)
    wk = np.full(len(f), 2.0); wk[0] = 1.0; wk[-1] = 1.0  # 單邊頻譜：DC 與奈奎斯特不加倍
    norm = NFFT * float(np.sum(win.astype(np.float64) ** 2))
    masks = {name: (f >= lo) & (f < hi) for name, lo, hi in BANDS}
    out = {name: np.empty(n) for name, *_ in BANDS}
    out["total"] = np.empty(n)
    xf = x.astype(np.float32) / 32768.0
    for a in range(0, n, 2000):
        b = min(n, a + 2000)
        idx = (np.arange(a, b) * hop)[:, None] + np.arange(NFFT)[None, :]
        p = np.abs(np.fft.rfft(xf[idx] * win, axis=1)) ** 2 * wk / norm
        for name, m in masks.items():
            out[name][a:b] = p[:, m].sum(axis=1)
        out["total"][a:b] = p.sum(axis=1)
    t = np.arange(n) * HOP_S + NFFT / 2 / sr
    return t, out


def db(ms):
    return 10 * np.log10(np.asarray(ms, dtype=np.float64) + 1e-12)


def mean_db(ms, sel):
    return float(db(ms[sel].mean())) if sel.any() else np.nan


def load(sess, raw, cache):
    """{channel: (t, {band: ms})}, cached per session."""
    out = {}
    for ch in ["s0", "s1", "s2"]:
        p = find(os.path.join(raw, sess.dir), f"_{ch}.wav")
        if not p:
            continue
        cf = os.path.join(cache, f"{sess.session_id}_{ch}.npz")
        if os.path.exists(cf):
            z = np.load(cf)
            out[ch] = (z["t"], {k: z[k] for k in z.files if k != "t"})
        else:
            x, sr = read_wav_int16(p)
            t, ms = band_ms(x, sr)
            np.savez(cf, t=t, **ms)
            out[ch] = (t, ms)
    return out


def plot(sess, data, cut, ctx, path):
    chans = list(data)
    fig, axes = plt.subplots(len(chans), 1, figsize=(15, 2.4 * len(chans) + 0.9), sharex=True, squeeze=False)
    for ax, ch in zip(axes[:, 0], chans):
        t, ms = data[ch]
        k = int(round(1 / HOP_S))
        nb = len(t) // k
        tt = t[:nb * k].reshape(nb, k).mean(1)
        lo = []
        for name, *_ in BANDS:
            y = db(ms[name][:nb * k].reshape(nb, k).mean(1))
            ax.plot(tt, y, lw=0.7, color=COLORS[name], label=f"{name} Hz")
            lo.append(np.percentile(y, 1))
        if sess.domain == "lab":
            for r in cut.itertuples():
                ax.axvspan(r.start, r.end, color="0.88", lw=0, zorder=0)
        else:
            c = ctx.sort_values("t")
            run = c.spindle_speed_rpm.fillna(0).values > 0
            te = np.append(c.t.values, c.t.values[-1] + 0.1)
            for i in np.flatnonzero(run):
                ax.axvspan(te[i], te[i + 1], ymin=0, ymax=0.04, color="0.45", lw=0, zorder=0)
        ax.set_ylim(min(lo) - 5, 2)
        ax.set_ylabel(f"{CH_NAME[ch]}\ndB（相對滿刻度）", fontsize=8)
        ax.grid(axis="y", color="0.92", lw=0.5)
    axes[0, 0].legend(loc="upper right", fontsize=8, ncol=4, frameon=False)
    axes[-1, 0].set_xlabel("時間（秒）；每秒能量平均")
    axes[-1, 0].set_xlim(0, sess.s0_sec)
    shade = "灰底＝切削路徑" if sess.domain == "lab" else "底部深灰細帶＝主軸運轉（控制器）"
    fig.suptitle(f"{SHORT[sess.machine]}  dataset{sess.dataset}  ({sess.session_id})  {sess.s0_sec/60:.1f} 分鐘  ·  {shade}",
                 fontsize=9, x=0.01, ha="left")
    fig.tight_layout(rect=(0, 0, 1, 0.96))
    fig.savefig(path, dpi=72)
    plt.close(fig)


def main():
    db_dir, raw, out = sys.argv[1], sys.argv[2], sys.argv[3]
    cache, pngd = os.path.join(out, "cache"), os.path.join(out, "png")
    os.makedirs(cache, exist_ok=True); os.makedirs(pngd, exist_ok=True)
    m = sqlite3.connect(os.path.join(db_dir, "main.sqlite"))
    ses = pd.read_sql("select * from sessions", m).merge(
        pd.read_sql("select session_id, dataset, dir from session_source", sqlite3.connect(os.path.join(db_dir, "source", "source.sqlite"))),
        on="session_id").sort_values(["machine", "dataset"]).reset_index(drop=True)
    ev = pd.read_sql("select * from events where cut_state='cutting'", m)
    ctx_all = pd.read_sql("select session_id, t, spindle_speed_rpm from context_ts", m)

    files, paths, segs, trend, pngs, energy = [], [], [], [], [], []
    for s in ses.itertuples():
        data = load(s, raw, cache)
        cut = ev[ev.session_id == s.session_id].sort_values("start")
        ctx = ctx_all[ctx_all.session_id == s.session_id]
        is_test_line = s.domain == "line" and s.dataset not in VF10_REF
        base = dict(機台=SHORT[s.machine], 場次=f"ds{s.dataset:02d}", session_id=s.session_id)

        for ch, (t, ms) in data.items():
            row = dict(base, 聲道=CH_NAME[ch], 秒數=round(float(t[-1] + HOP_S / 2), 1), 冷卻=s.coolant,
                       VF10角色=("測試" if is_test_line else "參考") if s.domain == "line" else "")
            row.update({f"{b} dB": round(mean_db(ms[b], np.ones(len(t), bool)), 2) for b, *_ in BANDS})
            files.append(row)
            energy.append(np.max(np.abs(db(sum(ms[b] for b in PARTITION)) - db(ms["total"]))))
            k = int(round(1 / HOP_S)); nb = len(t) // k
            tb = pd.DataFrame({"秒": np.arange(nb)})
            for b, *_ in BANDS:
                tb[f"{b} dB"] = np.round(db(ms[b][:nb * k].reshape(nb, k).mean(1)), 2)
            trend.append(tb.assign(**{k_: v for k_, v in base.items()}, 聲道=ch))

        if not is_test_line:  # 實驗室全部路徑＋VF-10 參考場次路徑
            for r in cut.itertuples():
                row = dict(base, 路徑序=int(r.idx), 起秒=r.start, 迄秒=r.end, 長度秒=r.dur, 材料=r.material, 刀具類別=r.tool_class,
                           N_rpm=r.N, F_mm_s=r.F, ap_mm=r.ap, ae_mm=r.ae, Z刃=r.Z, D_mm=r.D, 方向=r.direction, 冷卻=r.coolant)
                for ch, (t, ms) in data.items():
                    sel = (t >= r.start) & (t < r.end)
                    for b, *_ in BANDS:
                        row[f"{ch} {b} dB"] = round(mean_db(ms[b], sel), 2)
                paths.append(row)
        else:  # VF-10 測試場次：固定 10 秒段，只帶控制器資訊與場次層級工況
            c = ctx.sort_values("t")
            for a in np.arange(0, s.s0_sec, SEG_S):
                b_ = min(a + SEG_S, s.s0_sec)
                w = c[(c.t >= a) & (c.t < b_)]
                spd = w.spindle_speed_rpm
                row = dict(base, 段起秒=round(a, 1), 段迄秒=round(b_, 1), 主軸運轉比例=round(float((spd.fillna(0) > 0).mean()) if len(w) else np.nan, 2),
                           主軸轉速中位_rpm=float(spd[spd > 0].median()) if (spd > 0).any() else 0.0,
                           場次材料=", ".join(sorted(cut.material.unique())), 場次刀具=", ".join(sorted(cut.tool_class.unique())),
                           場次N範圍=rng(cut.N), 場次F範圍=rng(cut.F))
                for ch, (t, ms) in data.items():
                    sel = (t >= a) & (t < b_)
                    for bn, *_ in BANDS:
                        row[f"{ch} {bn} dB"] = round(mean_db(ms[bn], sel), 2)
                segs.append(row)

        fn = f"{s.machine}_ds{s.dataset:02d}_bands.png"
        plot(s, data, cut, ctx, os.path.join(pngd, fn))
        pngs.append((s.machine, os.path.join(pngd, fn), f"{SHORT[s.machine]} dataset{s.dataset}"))
        print(fn, flush=True)

    P = pd.DataFrame(paths)
    s0c = [f"s0 {b} dB" for b, *_ in BANDS]
    coarse = P.groupby(["機台", "材料", "刀具類別", "方向"]).agg(路徑數=("路徑序", "size"), 總秒數=("長度秒", "sum"), 場次數=("場次", "nunique"),
                                                             **{f"s0 {b} 平均dB": (f"s0 {b} dB", "mean") for b, *_ in BANDS}).round(2).reset_index()
    keys = ["機台", "材料", "刀具類別", "N_rpm", "F_mm_s", "ap_mm", "ae_mm", "Z刃", "D_mm", "方向"]
    fine = P.groupby(keys).agg(路徑數=("路徑序", "size"), 總秒數=("長度秒", "sum"), 場次數=("場次", "nunique"),
                               **{f"s0 {b} 平均dB": (f"s0 {b} dB", "mean") for b, *_ in BANDS}).round(2).reset_index()
    T = pd.concat(trend, ignore_index=True)[["機台", "場次", "session_id", "聲道", "秒"] + [f"{b} dB" for b, *_ in BANDS]]
    F = pd.DataFrame(files)
    S = pd.DataFrame(segs)

    evs = ev.merge(ses[["session_id", "has_s1", "has_s2"]], on="session_id")
    n_files, n_paths_all, n_pc = len(F), len(ev), int(len(evs) + evs.has_s1.sum() + evs.has_s2.sum())
    readme = pd.DataFrame([
        ("音檔", n_files, "VM20i 111（s0、s1、s2 各 37）、VMX30Ui 38（s0、s2 各 19）、VF-10 26（s0、s2 各 13）"),
        ("錄音場次", len(ses), f"共 {ses.s0_sec.sum()/3600:.2f} 小時；每場同時錄 2 到 3 個聲道"),
        ("切削路徑", n_paths_all, "VM20i 1,585、VMX30Ui 927、VF-10 63；工況記在每條路徑上"),
        ("路徑×聲道", n_pc, "切削路徑乘上該場的聲道數"),
        ("本表『切削路徑』頁", len(P), "實驗室全部路徑＋VF-10 參考場次（ds2、ds7、ds12）的路徑"),
        ("本表『VF-10測試場_10秒段』頁", len(S), "VF-10 其餘 10 個測試場次，每 10 秒一段，不列標註區間（封存）"),
    ], columns=["項目", "數量", "說明"])
    notes = pd.DataFrame([
        ("頻帶", "0–4k、1–4k、4–10k、10–24k Hz。0–4k 包含 1–4k；0–4k＋4–10k＋10–24k＝全頻。"),
        ("dB", "相對滿刻度（同 dBFS 基準）。麥克風未校正，不是聲壓 dB SPL，只能做相對比較。"),
        ("計算", "每 0.1 秒一格，4096 點 Hann 窗 FFT（11.7 Hz 解析），Parseval 換算各頻帶均方；路徑、段、每秒的值都是能量平均後取 dB。"),
        ("工況單位", "N＝rpm、F＝mm/s（已對照控制器確認）、ap／ae／D＝mm、Z＝刃數；方向 DOWN＝順銑、UP＝逆銑；刀具類別已去缺陷星號。"),
        ("不含標籤", "本檔不含正常／異常標籤。VF-10 測試場次的標註區間會洩漏異常位置，依 2026-10-14 預先登記封存，所以改用 10 秒段。"),
        ("能量守恆驗證", f"全部 {n_files} 個音檔、每一格：三段頻帶相加與全頻的最大差 {max(energy):.2e} dB。"),
    ], columns=["項目", "內容"])

    xlsx = os.path.join(out, "MSM_頻帶趨勢與工況.xlsx")
    with pd.ExcelWriter(xlsx, engine="openpyxl") as w:
        readme.to_excel(w, sheet_name="說明", index=False, startrow=0)
        notes.to_excel(w, sheet_name="說明", index=False, startrow=len(readme) + 3)
        F.to_excel(w, sheet_name="音檔清單", index=False)
        coarse.to_excel(w, sheet_name="工況分類_粗", index=False)
        fine.to_excel(w, sheet_name="工況分類_細", index=False)
        P.to_excel(w, sheet_name="切削路徑", index=False)
        S.to_excel(w, sheet_name="VF-10測試場_10秒段", index=False)
        T.to_excel(w, sheet_name="每秒趨勢", index=False)
    wb = load_workbook(xlsx)
    head = PatternFill("solid", fgColor="DCE6EC")
    for ws in wb.worksheets:
        ws.freeze_panes = "A2" if ws.title != "說明" else None
        for cell in ws[1]:
            cell.font = Font(bold=True); cell.fill = head; cell.alignment = Alignment(wrap_text=True, vertical="top")
        for col in ws.columns:
            width = max(len(str(c.value)) if c.value is not None else 0 for c in list(col)[:200])
            ws.column_dimensions[col[0].column_letter].width = min(60, max(8, width * 1.6 if any(ord(ch) > 255 for ch in str(col[0].value)) else width + 2))
    hdr2 = len(readme) + 4
    for cell in wb["說明"][hdr2]:
        cell.font = Font(bold=True); cell.fill = head
    for mach in ["vm20i", "vmx30ui", "vf10"]:
        ws = wb.create_sheet(f"趨勢圖_{SHORT[mach]}")
        ws["A1"] = f"{SHORT[mach]}：每場一張，每個聲道一列，四條線＝四個頻帶（每秒能量平均）"
        ws["A1"].font = Font(bold=True)
        r = 3
        for mm, p, cap in pngs:
            if mm != mach:
                continue
            img = XLImage(p)
            img.width, img.height = img.width * 0.85, img.height * 0.85
            ws.add_image(img, f"A{r}")
            r += int(img.height / 20) + 2
    wb.save(xlsx)
    print(f"saved {xlsx}; files {n_files}, paths rows {len(P)}, vf10 test segs {len(S)}, trend rows {len(T)}, max partition err {max(energy):.2e} dB")


if __name__ == "__main__":
    main()
