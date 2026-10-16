"""Per-audio-file report: machine, conditions, and a dB-vs-seconds chart for each of the 175 recordings.

Spec: design/bdd-band-trends.md (v2).  Usage: python per_file_report.py <db_dir> <raw_dir> <band_out_dir>
Reads the band cache written by band_trends.py (<band_out_dir>/cache); writes <band_out_dir>/per_file_png/*.png
and <band_out_dir>/MSM_每音檔聲音與工況.xlsx. No labels; VF-10 test sessions use 10-s segments only.
"""
import os
import sqlite3
import sys

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill

from band_trends import BANDS, CH_NAME, COLORS, HOP_S, SEG_S, SHORT, VF10_REF, db, load, mean_db
from plot_db_envelope import rng

TYPE = {"vm20i": "Hurco VM20i 三軸立式銑床", "vmx30ui": "Hurco VMX30Ui 五軸銑床", "vf10": "Haas VF-10 產線立式銑床"}
SITE = {"lab": "實驗室（Purdue IMI）", "line": "產線"}
MIC = {"s0": "空氣麥克風 Fifine K053", "s1": "接觸式聽診器 1", "s2": "接觸式聽診器 2"}
HEAD = PatternFill("solid", fgColor="DCE6EC")
LEVELS = [("全頻", "total")] + [(b, b) for b, *_ in BANDS]


def plot_file(s, ch, t, ms, cut, ctx, path):
    k = int(round(1 / HOP_S)); nb = len(t) // k
    tt = t[:nb * k].reshape(nb, k).mean(1)
    fig, ax = plt.subplots(figsize=(13.5, 3.4))
    lows = []
    for name, key in LEVELS:
        y = db(ms[key][:nb * k].reshape(nb, k).mean(1))
        ax.plot(tt, y, lw=1.0 if key == "total" else 0.7, color="#1c2321" if key == "total" else COLORS[key],
                label="全頻 0-24k Hz" if key == "total" else f"{name} Hz", zorder=3 if key == "total" else 2)
        lows.append(np.percentile(y, 1))
    if s.domain == "lab":
        for r in cut.itertuples():
            ax.axvspan(r.start, r.end, color="0.88", lw=0, zorder=0)
    else:
        c = ctx.sort_values("t")
        run = c.spindle_speed_rpm.fillna(0).values > 0
        te = np.append(c.t.values, c.t.values[-1] + 0.1)
        for i in np.flatnonzero(run):
            ax.axvspan(te[i], te[i + 1], ymin=0, ymax=0.04, color="0.45", lw=0, zorder=0)
    ax.set_ylim(min(lows) - 5, 2); ax.set_xlim(0, s.s0_sec)
    ax.set_xlabel("時間（秒）"); ax.set_ylabel("dB（相對滿刻度）")
    ax.grid(axis="y", color="0.92", lw=0.5)
    ax.legend(loc="lower right", fontsize=7, ncol=5, frameon=False)
    shade = "灰底＝切削路徑" if s.domain == "lab" else "底部深灰帶＝主軸運轉"
    ax.set_title(f"{SHORT[s.machine]} dataset{s.dataset} · {CH_NAME[ch]} · {shade}", fontsize=9, loc="left")
    fig.tight_layout(); fig.savefig(path, dpi=72); plt.close(fig)


def write_table(ws, df, row, col=1):
    for j, c in enumerate(df.columns):
        cell = ws.cell(row=row, column=col + j, value=c)
        cell.font = Font(bold=True); cell.fill = HEAD; cell.alignment = Alignment(wrap_text=True, vertical="top")
    for i, rec in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(rec):
            ws.cell(row=row + i, column=col + j, value=None if (isinstance(v, float) and np.isnan(v)) else v)
    return row + len(df) + 1


def main():
    db_dir, raw, out = sys.argv[1], sys.argv[2], sys.argv[3]
    cache, pngd = os.path.join(out, "cache"), os.path.join(out, "per_file_png")
    os.makedirs(pngd, exist_ok=True)
    m = sqlite3.connect(os.path.join(db_dir, "main.sqlite"))
    ses = pd.read_sql("select * from sessions", m).merge(
        pd.read_sql("select session_id, dataset, dir from session_source", sqlite3.connect(os.path.join(db_dir, "source", "source.sqlite"))),
        on="session_id").sort_values(["machine", "dataset"]).reset_index(drop=True)
    ev = pd.read_sql("select * from events where cut_state='cutting'", m)
    ctx_all = pd.read_sql("select session_id, t, spindle_speed_rpm from context_ts", m)

    blocks = {k: [] for k in SHORT}  # machine -> [(summary_row, info, tl_title, tl, png)]
    for s in ses.itertuples():
        data = load(s, raw, cache)
        cut = ev[ev.session_id == s.session_id].sort_values("start")
        ctx = ctx_all[ctx_all.session_id == s.session_id]
        test_line = s.domain == "line" and s.dataset not in VF10_REF
        for ch, (t, ms) in data.items():
            name_ = f"{SHORT[s.machine]}_ds{s.dataset:02d}_{ch}"
            png = os.path.join(pngd, f"{name_}.png")
            if not os.path.exists(png):
                plot_file(s, ch, t, ms, cut, ctx, png)
            allsel = np.ones(len(t), bool)
            lv = {name: round(mean_db(ms[key], allsel), 2) for name, key in LEVELS}
            cond = dict(材料=", ".join(sorted(cut.material.unique())), 刀具類別=", ".join(sorted(cut.tool_class.unique())),
                        N=rng(cut.N), F=rng(cut.F), ap=rng(cut.ap), ae=rng(cut.ae), Z=rng(cut.Z), D=rng(cut.D),
                        方向="/".join(sorted(cut.direction.unique())).replace("DOWN", "順銑").replace("UP", "逆銑"))
            info = [("機台", f"{SHORT[s.machine]}（{TYPE[s.machine]}）"), ("場域／冷卻", f"{SITE[s.domain]}，{'濕切' if s.coolant == 'wet' else '乾切'}"),
                    ("場次", f"dataset{s.dataset}（{s.session_id}）"), ("聲道", f"{ch}：{MIC[ch]}"), ("長度", f"{s.s0_sec:.1f} 秒（{s.s0_sec/60:.1f} 分）"),
                    ("切削路徑數", len(cut)), ("材料", cond["材料"]), ("刀具類別", cond["刀具類別"]), ("主軸轉速 N", f"{cond['N']} rpm"),
                    ("進給 F", f"{cond['F']} mm/s"), ("軸向切深 ap", f"{cond['ap']} mm"), ("徑向切寬 ae", f"{cond['ae']} mm"),
                    ("刃數 Z", cond["Z"]), ("刀徑 D", f"{cond['D']} mm"), ("方向", cond["方向"])] + \
                   [(f"整檔平均 {name} dB", v) for name, v in lv.items()]
            if not test_line:
                rows = []
                for r in cut.itertuples():
                    sel = (t >= r.start) & (t < r.end)
                    row = {"路徑序": int(r.idx), "起秒": r.start, "迄秒": r.end, "長度秒": r.dur, "材料": r.material, "刀具類別": r.tool_class,
                           "N rpm": r.N, "F mm/s": r.F, "ap mm": r.ap, "ae mm": r.ae, "Z 刃": r.Z, "D mm": r.D,
                           "方向": "順銑" if r.direction == "DOWN" else "逆銑"}
                    row.update({f"{name} dB": round(mean_db(ms[key], sel), 2) for name, key in LEVELS})
                    rows.append(row)
                tl, tl_title = pd.DataFrame(rows), "時間軸工況表（每條切削路徑；對照圖中灰底）"
            else:
                c = ctx.sort_values("t"); rows = []
                for a in np.arange(0, s.s0_sec, SEG_S):
                    b_ = min(a + SEG_S, s.s0_sec)
                    sp = c[(c.t >= a) & (c.t < b_)].spindle_speed_rpm
                    sel = (t >= a) & (t < b_)
                    row = {"段起秒": round(a, 1), "段迄秒": round(b_, 1), "主軸運轉比例": round(float((sp.fillna(0) > 0).mean()), 2) if len(sp) else None,
                           "主軸轉速中位 rpm": float(sp[sp > 0].median()) if (sp > 0).any() else 0.0}
                    row.update({f"{name} dB": round(mean_db(ms[key], sel), 2) for name, key in LEVELS})
                    rows.append(row)
                tl = pd.DataFrame(rows)
                tl_title = "時間軸表（每 10 秒一段；VF-10 測試場次不列標註區間，依預先登記封存）"
            summ = {"音檔": f"dataset{s.dataset} {ch}", "場次": f"ds{s.dataset:02d}", "session_id": s.session_id, "聲道": ch, "麥克風": MIC[ch],
                    "冷卻": "濕切" if s.coolant == "wet" else "乾切", "長度秒": round(s.s0_sec, 1), "路徑數": len(cut),
                    "時間軸表": "10 秒段（封存）" if test_line else "逐條路徑", "材料": cond["材料"], "刀具類別": cond["刀具類別"],
                    "N rpm": cond["N"], "F mm/s": cond["F"], "ap mm": cond["ap"], "ae mm": cond["ae"], "Z 刃": cond["Z"], "D mm": cond["D"],
                    "方向": cond["方向"], **{f"{name} dB": v for name, v in lv.items()}}
            blocks[s.machine].append((summ, info, tl_title, tl, png))
            print(name_, flush=True)

    wb = Workbook(); wb.remove(wb.active)
    total = 0
    for mach in ["vm20i", "vmx30ui", "vf10"]:
        bl = blocks[mach]; total += len(bl)
        ws = wb.create_sheet(SHORT[mach])
        ws["A1"] = f"{SHORT[mach]}（{TYPE[mach]}）：{len(bl)} 個音檔"; ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = ("dB 為相對滿刻度（麥克風未校正，只能相對比較），整檔平均為能量平均。工況單位：N rpm、F mm/s、ap／ae／D mm。"
                    "不含正常／異常標籤。點「音檔」欄可跳到下方該音檔的區塊。")
        if mach == "vf10":
            ws["A3"] = "VF-10：ds2、ds7、ds12（參考場次）逐條路徑；其餘 10 個測試場次只列 10 秒段，不列標註區間（依 2026-10-14 預先登記封存）。"
        S = pd.DataFrame([b[0] for b in bl])
        top = 5
        after = write_table(ws, S, top)
        row = after + 3
        for i, (summ, info, tl_title, tl, png) in enumerate(bl):
            anchor = row
            ws.cell(row=row, column=1, value=f"▍{SHORT[mach]} {summ['音檔']}（{summ['麥克風']}）").font = Font(bold=True, size=13, color="2B5C7A")
            back = ws.cell(row=row, column=4, value="↑ 回總表"); back.hyperlink = f"#'{SHORT[mach]}'!A{top}"
            back.font = Font(color="2B5C7A", underline="single")
            for j, (k_, v) in enumerate(info, start=1):
                ws.cell(row=row + j, column=1, value=k_).font = Font(bold=True)
                ws.cell(row=row + j, column=2, value=v)
            img = XLImage(png); ws.add_image(img, f"D{row + 1}")
            r2 = row + max(len(info), int(img.height / 20) + 1) + 2
            ws.cell(row=r2, column=1, value=tl_title).font = Font(bold=True)
            row = write_table(ws, tl, r2 + 1) + 3
            link = ws.cell(row=top + 1 + i, column=1)
            link.hyperlink = f"#'{SHORT[mach]}'!A{anchor}"; link.font = Font(color="2B5C7A", underline="single")
        for col, w in zip("ABCDEFGHIJKLMNOPQRSTUVW", [22, 40, 18, 9, 22, 6, 9, 8, 14, 26, 14, 12, 14, 12, 12, 8, 10, 8, 9, 9, 9, 9, 9]):
            ws.column_dimensions[col].width = w
    xlsx = os.path.join(out, "MSM_每音檔聲音與工況.xlsx")
    wb.save(xlsx)
    print(f"saved {xlsx}: {total} files, sheets {wb.sheetnames}")

if __name__ == "__main__":
    main()
